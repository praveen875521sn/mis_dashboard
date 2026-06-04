"""
Management command: import_worksetu

Loads Supply_WorkSetu_Sub_Cluster.xlsx into WorkSetuPINMap.

Columns expected (4):
  Supply State | Supply District | Supply PIN | Sub Cluster

Usage:
  python manage.py import_worksetu
  python manage.py import_worksetu --file=data/Supply_WorkSetu_Sub_Cluster.xlsx
"""

import os
import openpyxl
from django.core.management.base import BaseCommand
from dashboard.models import WorkSetuPINMap


class Command(BaseCommand):
    help = 'Import WorkSetu Supply PIN → Sub Cluster mapping'

    def add_arguments(self, parser):
        parser.add_argument(
            '--file',
            default=os.path.join('data', 'Supply_WorkSetu_Sub_Cluster.xlsx'),
            help='Path to Supply_WorkSetu_Sub_Cluster.xlsx',
        )

    def handle(self, *args, **options):
        path = options['file']
        self.stdout.write(f'Importing WorkSetu PIN map from {path} ...')

        if not os.path.exists(path):
            self.stdout.write(self.style.WARNING(
                f'  File not found: {path} — skipping.'
            ))
            return

        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        ws = wb.active

        headers = [str(c.value).strip() if c.value else '' for c in next(ws.iter_rows(max_row=1))]

        def col(row, name):
            try:
                return row[headers.index(name)].value
            except (ValueError, IndexError):
                return None

        def _s(v):
            if v is None:
                return ''
            if isinstance(v, float):
                return str(int(v))
            return str(v).strip()

        self.stdout.write('  Clearing existing WorkSetu PIN data...')
        WorkSetuPINMap.objects.all().delete()

        batch = []
        skipped = 0
        BATCH_SIZE = 500

        for row in ws.iter_rows(min_row=2):
            pin = _s(col(row, 'Supply PIN'))
            sc  = _s(col(row, 'Sub Cluster'))
            if not pin or not sc:
                skipped += 1
                continue
            batch.append(WorkSetuPINMap(
                supply_state         = _s(col(row, 'Supply State')),
                supply_district      = _s(col(row, 'Supply District')),
                supply_pin           = pin,
                worksetu_sub_cluster = sc,
            ))
            if len(batch) >= BATCH_SIZE:
                WorkSetuPINMap.objects.bulk_create(batch, ignore_conflicts=True)
                batch = []

        if batch:
            WorkSetuPINMap.objects.bulk_create(batch, ignore_conflicts=True)

        total = WorkSetuPINMap.objects.count()
        self.stdout.write(self.style.SUCCESS(
            f'  ✓ {total} WorkSetu PIN→SubCluster mappings loaded ({skipped} skipped).'
        ))
