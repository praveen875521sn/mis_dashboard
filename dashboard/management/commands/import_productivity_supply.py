"""
Management command: import_productivity_supply
Loads the four new datasets that replace / augment the SAHI module:

    • Trainer_productivity_present.xlsx       → TrainerProductivityPresent(+Day)
    • Trainer_productivity_duration.xlsx      → TrainerProductivityDuration(+Day)
    • Associate_data.xlsx                     → AssociateData
    • ITI_Polytechnic_Colleges_by_PIN_Supply  → ITIPolytechnicByPIN

These replace the Trainer_Target.xlsx import (which is now removed) and add
the new SAHI "Supply Chain" lookups driven by Sub Cluster ID + Supply PIN.

Usage:
  python manage.py import_productivity_supply
  python manage.py import_productivity_supply --data-dir data
  python manage.py import_productivity_supply --only present
"""

import os
from datetime import date, datetime, time, timedelta

import openpyxl
from django.core.management.base import BaseCommand

from dashboard.models import (
    TrainerProductivityPresent, TrainerProductivityPresentDay,
    TrainerProductivityDuration, TrainerProductivityDurationDay,
    AssociateData, ITIPolytechnicByPIN,
)


# ─── Tiny coercers ──────────────────────────────────────────────────────────

def _s(v):
    return '' if v is None else str(v).strip()


def _i(v):
    try:
        if v is None or v == '':
            return 0
        return int(float(v))
    except (TypeError, ValueError):
        return 0


def _i_or_none(v):
    if v is None or v == '':
        return None
    try:
        return int(float(v))
    except (TypeError, ValueError):
        return None


def _f(v):
    """Cell → decimal hours/float. Accepts numbers, datetime.time, timedelta."""
    if v is None or v == '':
        return 0.0
    if isinstance(v, time):
        return v.hour + (v.minute / 60.0) + (v.second / 3600.0)
    if isinstance(v, timedelta):
        return v.total_seconds() / 3600.0
    if isinstance(v, datetime):
        return v.hour + (v.minute / 60.0) + (v.second / 3600.0)
    try:
        if isinstance(v, float) and (v != v):  # NaN
            return 0.0
        return float(v)
    except (TypeError, ValueError):
        return 0.0


def _f_or_none(v):
    if v is None or v == '':
        return None
    try:
        x = float(v)
        if x != x:  # NaN
            return None
        return x
    except (TypeError, ValueError):
        return None


def _pin(v):
    """PINs come as numbers, strings, or floats — normalise to 6-char string."""
    if v is None or v == '':
        return ''
    try:
        n = int(float(v))
        return str(n)
    except (TypeError, ValueError):
        return _s(v)


def _norm(name):
    return ' '.join(_s(name).lower().split())


def _idx(headers, *candidates):
    norm = [_norm(h) for h in headers]
    for c in candidates:
        cn = _norm(c)
        for i, h in enumerate(norm):
            if h == cn:
                return i
    return None


# ─── Command ────────────────────────────────────────────────────────────────

class Command(BaseCommand):
    help = ('Import Trainer productivity (present + duration) and SAHI supply-chain '
            '(Associate data + ITI/Polytechnic by PIN) data.')

    def add_arguments(self, parser):
        parser.add_argument('--data-dir', default='data')
        parser.add_argument(
            '--only',
            choices=['present', 'duration', 'associate', 'iti'],
            help='Run only one importer (default: all four).',
        )

    def handle(self, *args, **options):
        data_dir = options['data_dir']
        only     = options['only']

        plan = [
            ('present',   'Trainer_productivity_present.xlsx',
                self._import_present),
            ('duration',  'Trainer_productivity_duration.xlsx',
                self._import_duration),
            ('associate', 'Associate_data.xlsx',
                self._import_associate),
            ('iti',       'ITI_Polytechnic_Colleges_by_PIN_Supply.xlsx',
                self._import_iti),
        ]

        for key, fname, fn in plan:
            if only and only != key:
                continue
            path = os.path.join(data_dir, fname)
            if not os.path.exists(path):
                self.stdout.write(self.style.WARNING(
                    f'  ({fname} not found in {data_dir}/) — skipped'))
                continue
            fn(path)

        self.stdout.write(self.style.SUCCESS('\n✓ Productivity / supply imports complete.'))

    # ── Trainer productivity (present count) ────────────────────────────────
    def _import_present(self, path):
        self.stdout.write(f'Importing Trainer Productivity (Present) from {path} ...')
        wb = openpyxl.load_workbook(path, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            self.stdout.write('  Empty — skipped.'); return

        headers = rows[0]
        ix = {
            'batch':     _idx(headers, 'Batch ID'),
            'centre':    _idx(headers, 'Centre Name', 'Center Name'),
            'centre_id': _idx(headers, 'Centre ID', 'Center ID'),
            'qp':        _idx(headers, 'QP Name', 'QP'),
            'trainer':   _idx(headers, 'Trainer'),
            'tid':       _idx(headers, 'Trainer ID'),
            'enrolled':  _idx(headers, 'Enrolled'),
        }
        # Day columns: header is a datetime/date
        day_cols = []
        for i, h in enumerate(headers):
            if isinstance(h, datetime):
                day_cols.append((i, h.date()))
            elif isinstance(h, date):
                day_cols.append((i, h))

        TrainerProductivityPresent.objects.all().delete()
        created = 0
        day_created = 0
        for row in rows[1:]:
            if not any(row):
                continue
            tid = _i(row[ix['tid']] if ix['tid'] is not None else None)
            if not tid:
                continue
            obj = TrainerProductivityPresent.objects.create(
                batch_id     = _s(row[ix['batch']])     if ix['batch']     is not None else '',
                centre_name  = _s(row[ix['centre']])    if ix['centre']    is not None else '',
                centre_id    = _s(row[ix['centre_id']]) if ix['centre_id'] is not None else '',
                qp_name      = _s(row[ix['qp']])        if ix['qp']        is not None else '',
                trainer_name = _s(row[ix['trainer']])   if ix['trainer']   is not None else '',
                trainer_id   = tid,
                enrolled     = _i(row[ix['enrolled']])  if ix['enrolled']  is not None else 0,
            )
            created += 1
            day_batch = []
            for col_idx, d in day_cols:
                val = row[col_idx] if col_idx < len(row) else None
                if val is None or val == '':
                    continue   # skip blanks — only store actual present counts
                day_batch.append(TrainerProductivityPresentDay(
                    parent=obj, date=d, present=_i(val)
                ))
            if day_batch:
                TrainerProductivityPresentDay.objects.bulk_create(day_batch)
                day_created += len(day_batch)

        self.stdout.write(f'  ✓ {created} present rows, {day_created} day-present entries.')

    # ── Trainer productivity (duration / hours) ─────────────────────────────
    def _import_duration(self, path):
        self.stdout.write(f'Importing Trainer Productivity (Duration) from {path} ...')
        wb = openpyxl.load_workbook(path, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            self.stdout.write('  Empty — skipped.'); return

        headers = rows[0]
        ix = {
            'batch':     _idx(headers, 'Batch ID'),
            'centre':    _idx(headers, 'Centre Name', 'Center Name'),
            'centre_id': _idx(headers, 'Centre ID', 'Center ID'),
            'qp':        _idx(headers, 'QP Name', 'QP'),
            'trainer':   _idx(headers, 'Trainer'),
            'tid':       _idx(headers, 'Trainer ID'),
        }
        day_cols = []
        for i, h in enumerate(headers):
            if isinstance(h, datetime):
                day_cols.append((i, h.date()))
            elif isinstance(h, date):
                day_cols.append((i, h))

        TrainerProductivityDuration.objects.all().delete()
        created = 0
        day_created = 0
        for row in rows[1:]:
            if not any(row):
                continue
            tid = _i(row[ix['tid']] if ix['tid'] is not None else None)
            if not tid:
                continue
            obj = TrainerProductivityDuration.objects.create(
                batch_id     = _s(row[ix['batch']])     if ix['batch']     is not None else '',
                centre_name  = _s(row[ix['centre']])    if ix['centre']    is not None else '',
                centre_id    = _s(row[ix['centre_id']]) if ix['centre_id'] is not None else '',
                qp_name      = _s(row[ix['qp']])        if ix['qp']        is not None else '',
                trainer_name = _s(row[ix['trainer']])   if ix['trainer']   is not None else '',
                trainer_id   = tid,
            )
            created += 1
            day_batch = []
            for col_idx, d in day_cols:
                val = row[col_idx] if col_idx < len(row) else None
                if val is None or val == '':
                    continue
                hours = _f(val)
                if hours <= 0:
                    continue
                day_batch.append(TrainerProductivityDurationDay(
                    parent=obj, date=d, hours=round(hours, 3)
                ))
            if day_batch:
                TrainerProductivityDurationDay.objects.bulk_create(day_batch)
                day_created += len(day_batch)

        self.stdout.write(f'  ✓ {created} duration rows, {day_created} day-hour entries.')

    # ── Associate Data ──────────────────────────────────────────────────────
    def _import_associate(self, path):
        self.stdout.write(f'Importing Associate Data from {path} ...')
        wb = openpyxl.load_workbook(path, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            self.stdout.write('  Empty — skipped.'); return

        headers = rows[0]
        ix = {
            'ano':      _idx(headers, 'Associate No', 'Associate ID'),
            'name':     _idx(headers, 'Associate', 'Associate Name'),
            'dob':      _idx(headers, 'Date of Birth', 'DOB'),
            'gender':   _idx(headers, 'Gender'),
            'addr':     _idx(headers, 'Permanent Address', 'Address'),
            'state':    _idx(headers, 'Supply State', 'State'),
            'district': _idx(headers, 'Supply District', 'District'),
            'pin':      _idx(headers, 'Supply PIN', 'PIN'),
            'lat':      _idx(headers, 'latitude', 'Latitude'),
            'lng':      _idx(headers, 'longitude', 'Longitude'),
            'dcity':    _idx(headers, 'Demand City'),
            'dcloc':    _idx(headers, 'Demand Cluster (Client Location)', 'Demand Cluster'),
            'client':   _idx(headers, 'Client'),
            'scid':     _idx(headers, 'Sub Cluster ID', 'Sub Project ID'),
            'cluster':  _idx(headers, 'Cluster'),
            'sc':       _idx(headers, 'Sub Cluster'),
        }

        def g(row, k):
            i = ix.get(k)
            return row[i] if (i is not None and i < len(row)) else None

        AssociateData.objects.all().delete()
        bulk = []
        for row in rows[1:]:
            if not any(row):
                continue
            bulk.append(AssociateData(
                associate_no               = _s(g(row, 'ano')),
                associate_name             = _s(g(row, 'name')),
                date_of_birth              = _s(g(row, 'dob')),
                gender                     = _s(g(row, 'gender')),
                permanent_address          = _s(g(row, 'addr')),
                supply_state               = _s(g(row, 'state')),
                supply_district            = _s(g(row, 'district')),
                supply_pin                 = _pin(g(row, 'pin')),
                latitude                   = _f_or_none(g(row, 'lat')),
                longitude                  = _f_or_none(g(row, 'lng')),
                demand_city                = _s(g(row, 'dcity')),
                demand_cluster_client_loc  = _s(g(row, 'dcloc')),
                client                     = _s(g(row, 'client')),
                sub_cluster_id             = _s(g(row, 'scid')),
                cluster                    = _s(g(row, 'cluster')),
                sub_cluster                = _s(g(row, 'sc')),
            ))
        AssociateData.objects.bulk_create(bulk, batch_size=2000)
        self.stdout.write(f'  ✓ {len(bulk)} associate rows.')

    # ── ITI / Polytechnic by PIN ────────────────────────────────────────────
    def _import_iti(self, path):
        self.stdout.write(f'Importing ITI / Polytechnic by PIN from {path} ...')
        wb = openpyxl.load_workbook(path, data_only=True)
        # The expected sheet name is "ITI & Polytechnic Colleges" — but fall
        # back to the first sheet if not present.
        ws = (wb['ITI & Polytechnic Colleges']
              if 'ITI & Polytechnic Colleges' in wb.sheetnames
              else wb.worksheets[0])
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            self.stdout.write('  Empty — skipped.'); return

        headers = rows[0]
        ix = {
            'sno':   _idx(headers, 'S.No.', 'S.No', 'SNo'),
            'pin':   _idx(headers, 'Supply PIN', 'PIN'),
            'lat':   _idx(headers, 'Latitude'),
            'lng':   _idx(headers, 'Longitude'),
            'inst':  _idx(headers, 'Institution Name', 'Name'),
            'type':  _idx(headers, 'Type'),
            'addr':  _idx(headers, 'Address / Location', 'Address', 'Location'),
            'dist':  _idx(headers, 'Approx. Distance (km)', 'Distance', 'Distance (km)'),
        }

        def g(row, k):
            i = ix.get(k)
            return row[i] if (i is not None and i < len(row)) else None

        ITIPolytechnicByPIN.objects.all().delete()
        bulk = []
        for row in rows[1:]:
            if not any(row):
                continue
            bulk.append(ITIPolytechnicByPIN(
                sno              = _i_or_none(g(row, 'sno')),
                supply_pin       = _pin(g(row, 'pin')),
                latitude         = _f_or_none(g(row, 'lat')),
                longitude        = _f_or_none(g(row, 'lng')),
                institution_name = _s(g(row, 'inst')),
                institution_type = _s(g(row, 'type')),
                address_location = _s(g(row, 'addr')),
                distance_km      = _s(g(row, 'dist')),
            ))
        ITIPolytechnicByPIN.objects.bulk_create(bulk, batch_size=1000)
        self.stdout.write(f'  ✓ {len(bulk)} ITI/Polytechnic rows.')
